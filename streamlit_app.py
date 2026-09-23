"""Preenchedor de Documentos PRONER — interface Streamlit."""

import io
import zipfile

import streamlit as st
from PIL import Image

from engine import (CADERNOS_DIR, DOCS, FIELD_IDS, FIELDS, build_memorial,
                    carregar_base_completa, compute_logo_extent_emu,
                    derive_instrumento_tokens, fill_template, parse_orcamento)

st.set_page_config(page_title="Preenchedor de Documentos PRONER",
                   page_icon="📋", layout="centered")

st.title("📋 Preenchedor de Documentos PRONER")
st.caption("Preencha uma vez — os documentos originais são preenchidos, sem alterar a formatação.")

# ── Logotipo ─────────────────────────────────────────────────────────
st.subheader("🏛️ Logotipo da Prefeitura")
logo_file = st.file_uploader("Imagem do brasão/logotipo", type=["png", "jpg", "jpeg", "webp", "bmp"])
logo_png_bytes, logo_cx, logo_cy = None, 900000, 900000
if logo_file is not None:
    _img = Image.open(logo_file).convert("RGBA")
    _buf = io.BytesIO()
    _img.save(_buf, format="PNG")
    logo_png_bytes = _buf.getvalue()
    logo_cx, logo_cy = compute_logo_extent_emu(_img.width, _img.height)
    st.image(logo_png_bytes, width=130, caption="Aparecerá centralizado no cabeçalho")

st.divider()

# ── Instrumento: Convênio ou Proposta ────────────────────────────────
st.subheader("📋 Instrumento")
c1, c2 = st.columns([1, 2])
tipo_instrumento = c1.radio("Tipo", ["Convênio", "Proposta"], index=1,
                            help="Em fase de proposta ainda não há número de convênio. "
                                 "Escolhendo 'Proposta', todos os documentos passam a dizer "
                                 "'proposta nº ...' no lugar de 'convênio nº ...'.")
numero_instrumento = c2.text_input(f"Nº da {tipo_instrumento} / Ano", placeholder="ex: 034959/2026")

vals = {}
vals["CONVENIO"] = numero_instrumento

c1, c2, c3 = st.columns([2, 2, 1])
vals["MUNICIPIO"] = c1.text_input("Município", placeholder="ex: João Lisboa")
vals["UF"] = c3.text_input("UF", placeholder="MA", max_chars=2)
vals["OBJETO"] = st.text_input("Objeto (como cadastrado no TransfereGov)",
                               placeholder="ex: Execução de obras e serviços de engenharia para estradas vicinais")
vals["DESCRICAO_OBJETO"] = st.text_area(
    "Descrição do Objeto", height=80,
    placeholder="ex: Contratação de obra civil para construção de rede para fomentar a geração e distribuição de energia")
vals["EIXO"] = st.text_input(
    "Eixo do objeto (IN 25)",
    placeholder="ex: Tópico 2 - Obras para o fomento ao acesso à energia elétrica")

# ── Prefeitura ───────────────────────────────────────────────────────
st.subheader("🏢 Dados da Prefeitura")
c1, c2 = st.columns(2)
vals["NOME_PREFEITO"] = c1.text_input("Nome do(a) Prefeito(a)")
vals["CNPJ_PREFEITURA"] = c2.text_input("CNPJ da Prefeitura", placeholder="00.000.000/0001-00")
vals["ENDERECO_SEDE"] = st.text_input("Endereço da Sede Administrativa")

# ── Obra ─────────────────────────────────────────────────────────────
st.subheader("🛣️ Dados da Obra")
vals["DESCRICAO_OBRA"] = st.text_input("Descrição da Obra")
vals["LOCALIZACAO_TRECHO"] = st.text_input("Localização / Trecho (texto corrido)")
vals["DATA_REFERENCIA"] = st.text_input("Data e local de referência",
                                        placeholder="ex: João Lisboa/MA, 23/09/2026")

# ── Trechos (tabela da Definição do Objeto) ──────────────────────────
st.subheader("📍 Trechos (Definição do Objeto)")
st.caption("Cada trecho vira uma linha na tabela 'Informações Básicas' do documento Definição do Objeto.")
qtd_trechos = st.number_input("Quantidade de trechos", min_value=0, max_value=30, value=1, step=1)

trechos = []
for i in range(int(qtd_trechos)):
    with st.expander(f"Trecho {i + 1}", expanded=(i == 0)):
        c1, c2 = st.columns([2, 1])
        nome = c1.text_input("Nome do trecho", value=f"Trecho {i + 1}", key=f"tr_nome_{i}")
        ext = c2.text_input("Extensão (km)", key=f"tr_ext_{i}", placeholder="ex: 12,5")
        c1, c2 = st.columns(2)
        ini = c1.text_input("Coordenadas iniciais", key=f"tr_ini_{i}",
                            placeholder="ex: 05°26'12\"S 47°24'30\"W")
        fim = c2.text_input("Coordenadas finais", key=f"tr_fim_{i}",
                            placeholder="ex: 05°28'44\"S 47°26'02\"W")
        trechos.append({"nome": nome, "coord_ini": ini, "coord_fim": fim, "extensao": ext})

# ── Responsável técnico ──────────────────────────────────────────────
st.subheader("👷 Responsável Técnico")
c1, c2 = st.columns(2)
vals["NOME_ENGENHEIRO"] = c1.text_input("Nome do Engenheiro")
vals["ESPECIALIDADE_ENG"] = c2.text_input("Especialidade", placeholder="ex: Engenheiro Civil")
c1, c2 = st.columns(2)
vals["CREA_ENGENHEIRO"] = c1.text_input("CREA nº", placeholder="ex: CREA/MA nº 12345D/MA")
vals["FUNCAO_ENG"] = c2.text_input("Função", placeholder="ex: Responsável técnico pela Fiscalização")

# ── Indicadores ──────────────────────────────────────────────────────
st.subheader("📊 Indicadores Municipais")
c1, c2, c3 = st.columns(3)
vals["POPULACAO_TOTAL"] = c1.text_input("População Total")
vals["POPULACAO_AFETADA"] = c2.text_input("População Afetada")
vals["RENDA_PER_CAPITA"] = c3.text_input("Renda Per Capita")
c1, c2, c3 = st.columns(3)
vals["IMPACTO_LOGISTICO"] = c1.text_input("Impacto Logístico")
vals["EXTENSAO_TOTAL"] = c2.text_input("Extensão Total de Estradas")
vals["PRECIPITACAO"] = c3.text_input("Precipitação Média Anual")

# ── Campos específicos ───────────────────────────────────────────────
st.subheader("🔴 Campos Específicos")
opcao_desoneracao = st.radio(
    "Regime de desoneração da folha", ["COM DESONERAÇÃO", "SEM DESONERAÇÃO"],
    horizontal=True,
    help="Preenche a Declaração de Desoneração de Folha com a alternativa adotada.")
c1, c2 = st.columns(2)
vals["LICENCA_JAZIDA"] = c1.text_input("Licença Ambiental da Jazida (protocolo)")
vals["LEGENDA_MAPA"] = c2.text_input("Legenda do Mapa/Fotos", placeholder="ex: MA")

# ── Memorial a partir da planilha ────────────────────────────────────
st.subheader("📑 Memorial Descritivo a partir da Planilha Orçamentária")

@st.cache_data(show_spinner="Indexando cadernos técnicos...")
def _base_cacheada(_versao: int = 0):
    return carregar_base_completa()


base_completa, indice = _base_cacheada(st.session_state.get("reindex", 0))
n_pdf = len(indice.get("servicos", {}))
n_total = len(base_completa.get("servicos", {}))
arquivos = indice.get("arquivos", [])

with st.expander(f"📚 Base de descrições técnicas — {n_total} serviço(s) disponível(is)"):
    if arquivos:
        st.write(f"**PDFs lidos de `base/cadernos/`:** {', '.join(arquivos)}")
        st.write(f"Serviços indexados automaticamente a partir dos PDFs: **{n_pdf}**")
    else:
        st.info("Nenhum PDF em `base/cadernos/`. Envie os cadernos técnicos do DNIT "
                "para essa pasta no GitHub e eles serão indexados automaticamente.")
    for aviso in indice.get("avisos", []):
        st.warning(aviso)

    if base_completa.get("servicos"):
        st.caption("Confira se cada código foi associado à seção correta do caderno:")
        linhas = [{"Código": c,
                   "Título extraído": d.get("titulo", ""),
                   "Seção": d.get("secao", "—"),
                   "Origem": d.get("fonte", "—"),
                   "Tamanho": f"{len(d.get('descricao',''))} car."}
                  for c, d in sorted(base_completa["servicos"].items())]
        st.dataframe(linhas, use_container_width=True, hide_index=True)
        cod = st.selectbox("Ver texto completo de um serviço",
                           [""] + sorted(base_completa["servicos"].keys()))
        if cod:
            st.text_area("Texto que será inserido no memorial",
                         base_completa["servicos"][cod].get("descricao", ""),
                         height=220)
    if st.button("🔄 Reindexar os PDFs"):
        carregar_base_completa(forcar_reindex=True)
        _base_cacheada.clear()
        st.session_state["reindex"] = st.session_state.get("reindex", 0) + 1
        st.rerun()

st.caption("Envie a planilha orçamentária (.xlsx ou .csv). O sistema identifica os códigos "
           "das composições e busca as descrições na base acima.")
planilha = st.file_uploader("Planilha orçamentária", type=["xlsx", "xls", "csv"],
                            key="orcamento")

memorial_texto = ""
if planilha is not None:
    try:
        rows = []
        if planilha.name.lower().endswith(".csv"):
            import csv
            text = planilha.getvalue().decode("utf-8", errors="replace")
            delim = ";" if text.count(";") > text.count(",") else ","
            rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(planilha.getvalue()), data_only=True)
            for sheet in wb.worksheets:
                for r in sheet.iter_rows(values_only=True):
                    rows.append(list(r))

        servicos = parse_orcamento(rows)
        if not servicos:
            st.warning("Nenhum serviço com código de composição foi identificado na planilha. "
                       "Confira se há uma coluna com o código (ex: 4015612).")
        else:
            memorial_texto, pendentes = build_memorial(servicos, base_completa)
            st.success(f"{len(servicos)} serviço(s) identificado(s) na planilha.")
            if pendentes:
                st.warning(
                    f"{len(pendentes)} serviço(s) sem descrição na base. Eles entram no "
                    f"memorial com a descrição da planilha e uma marcação para "
                    f"preenchimento manual:\n\n"
                    + "\n".join(f"- `{p['codigo']}` — {p['descricao'][:70]}" for p in pendentes))
            with st.expander("Pré-visualizar o memorial gerado"):
                st.text(memorial_texto[:4000] + ("..." if len(memorial_texto) > 4000 else ""))
    except Exception as e:
        st.error(f"Não consegui ler a planilha: {e}")

vals["MEMORIAL_SERVICOS"] = memorial_texto

st.divider()

# ── Seleção de documentos ────────────────────────────────────────────
st.subheader("📁 Documentos a Gerar")
marcar_todos = st.checkbox("Selecionar todos", value=True)
selecionados = []
cols = st.columns(2)
for i, (doc_id, label) in enumerate(DOCS):
    if cols[i % 2].checkbox(label, value=marcar_todos, key=f"doc_{doc_id}"):
        selecionados.append(doc_id)

st.divider()

# ── Geração ──────────────────────────────────────────────────────────
if st.button("⚡ Gerar Documentos", type="primary", use_container_width=True):
    if not selecionados:
        st.warning("Selecione pelo menos um documento.")
    else:
        # valores padrão para campos deixados em branco
        padroes = {f[0]: f[3] for f in FIELDS}
        municipio = vals.get("MUNICIPIO") or "XXXXXXXXXX"
        uf = vals.get("UF") or "XX"
        padroes["DATA_REFERENCIA"] = f"{municipio}/{uf}, XX/XX/XXXX"
        for fid in FIELD_IDS:
            if not vals.get(fid):
                vals[fid] = padroes.get(fid) or ""

        vals.update(derive_instrumento_tokens(tipo_instrumento, vals["CONVENIO"]))
        vals["OPCAO_DESONERACAO"] = opcao_desoneracao

        gerados = []
        barra = st.progress(0, text="Preenchendo...")
        for i, doc_id in enumerate(selecionados):
            label = dict(DOCS)[doc_id]
            barra.progress((i + 1) / len(selecionados), text=f"Preenchendo: {label}...")
            try:
                gerados.append((doc_id, label,
                                fill_template(doc_id, vals, logo_png_bytes,
                                              logo_cx, logo_cy, trechos)))
            except Exception as e:
                st.error(f"Erro em {label}: {e}")
        barra.empty()

        if gerados:
            st.success(f"{len(gerados)} documento(s) gerado(s).")
            for doc_id, label, data in gerados:
                st.download_button(
                    f"⬇ Baixar {label}", data=data, file_name=f"{doc_id}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"dl_{doc_id}", use_container_width=True)

            if len(gerados) > 1:
                zbuf = io.BytesIO()
                with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as zf:
                    for doc_id, _l, data in gerados:
                        zf.writestr(f"{doc_id}.docx", data)
                mun = (vals.get("MUNICIPIO") or "municipio").replace(" ", "_")
                num = (vals.get("CONVENIO") or "s-numero").replace("/", "-")
                st.download_button("📦 Baixar todos em ZIP", data=zbuf.getvalue(),
                                   file_name=f"DOCUMENTOS_{mun}_{num}.zip",
                                   mime="application/zip", use_container_width=True)

st.divider()
st.caption("Os documentos gerados são os arquivos originais do MAPA/PRONER, com a formatação "
           "preservada — apenas os campos são substituídos, e aparecem destacados em amarelo.")
